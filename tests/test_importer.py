"""Importer tests against a synthetic workbook shaped like the client's."""

from __future__ import annotations

import sqlite3

import pytest
from openpyxl import Workbook

from txtav import db, importer


@pytest.fixture
def workbook(tmp_path):
    def build(rows):
        wb = Workbook()
        ws = wb.active
        ws.title = importer.INPUT_SHEET
        ws.append(["PART NUMBER", "WH"])
        for row in rows:
            ws.append(list(row))
        wb.create_sheet(importer.OUTPUT_SHEET).append(
            [
                "PART NUMBER",
                "Alternate part number",
                "Description",
                "Condition code",
                "Quantity",
                "Outright price",
                "Sale price",
                "Core value",
                "Price reference",
            ]
        )
        path = tmp_path / "pn.xlsx"
        wb.save(path)
        return path

    return build


@pytest.fixture
def conn(tmp_path) -> sqlite3.Connection:
    connection = db.connect(tmp_path / "test.db")
    yield connection
    connection.close()


def test_imports_only_txt_rows(conn, workbook):
    path = workbook(
        [
            ("6622168-22", "TXT"),
            ("601R409101-341", "MHI"),
            ("2384029-8", "BBD"),
            ("260-1014-33", "TXT"),
        ]
    )
    report = importer.import_parts(conn, path)

    assert report.total_rows == 4
    assert report.txt_rows == 2
    assert report.imported == 2
    assert report.other_wh == {"MHI": 1, "BBD": 1}

    stored = {r["pn"] for r in conn.execute("SELECT pn FROM parts")}
    assert stored == {"6622168-22", "260-1014-33"}


def test_duplicates_are_skipped_case_insensitively(conn, workbook):
    path = workbook(
        [("MS21922-4c", "TXT"), ("MS21922-4C", "TXT"), ("6622168-22", "TXT")]
    )
    report = importer.import_parts(conn, path)

    assert report.imported == 2
    assert len(report.duplicates) == 1
    # The first spelling seen is the one kept and the one we will look up.
    assert conn.execute(
        "SELECT pn FROM parts WHERE dedupe_key = 'MS21922-4C'"
    ).fetchone()["pn"] == "MS21922-4c"


def test_blank_part_numbers_are_reported_not_imported(conn, workbook):
    path = workbook([("", "TXT"), ("   ", "TXT"), ("6622168-22", "TXT")])
    report = importer.import_parts(conn, path)

    assert report.imported == 1
    assert len(report.blank_pn_rows) == 2


def test_flags_part_numbers_needing_the_search_box(conn, workbook):
    path = workbook(
        [
            ("000-110105-52U/P", "TXT"),
            ("000-110105-79 U", "TXT"),
            ("6622168-22", "TXT"),
        ]
    )
    report = importer.import_parts(conn, path)

    # Only the slash needs the search box. The space survives percent encoding,
    # verified against the live portal on 2026-10-02.
    assert report.path_hostile == 1
    hostile = {
        r["pn"] for r in conn.execute("SELECT pn FROM parts WHERE path_hostile = 1")
    }
    assert hostile == {"000-110105-52U/P"}


def test_reimport_updates_rather_than_wiping(conn, workbook):
    # A cycle in progress has jobs and results pointing at these rows, so
    # deleting the table fails the foreign key check outright. An import is
    # also nearly always a refresh of the same list, per the scope.
    first = workbook([("6622168-22", "TXT"), ("260-1014-33", "TXT")])
    importer.import_parts(conn, first)
    second = workbook([("260-1014-33", "TXT"), ("9912125-2EX", "TXT")])
    report = importer.import_parts(conn, second)

    assert report.imported == 2
    stored = {r["pn"] for r in conn.execute("SELECT pn FROM parts")}
    assert stored == {"6622168-22", "260-1014-33", "9912125-2EX"}


def test_a_part_that_left_the_list_is_reported_not_deleted(conn, workbook):
    # Its results are still attached to it, and discarding a price someone may
    # be looking at is not a decision an import should take by itself.
    importer.import_parts(conn, workbook([("GONE-1", "TXT"), ("STAYS-1", "TXT")]))
    report = importer.import_parts(conn, workbook([("STAYS-1", "TXT")]))

    assert report.departed == ["GONE-1"]
    assert conn.execute(
        "SELECT COUNT(*) c FROM parts WHERE pn = 'GONE-1'"
    ).fetchone()["c"] == 1


def test_output_rows_are_consecutive_from_two(conn, workbook):
    # The output tab has a header at row 1 and one row per TXT part after it.
    # Source rows cannot serve: the real input tab carries 200 MHI and BBD rows
    # before the first TXT part.
    path = workbook(
        [("MHI-1", "MHI"), ("MHI-2", "MHI"), ("A-1", "TXT"), ("A-2", "TXT")]
    )
    importer.import_parts(conn, path)

    rows = list(conn.execute("SELECT pn, source_row, output_row FROM parts ORDER BY output_row"))
    assert [r["output_row"] for r in rows] == [2, 3]
    assert [r["source_row"] for r in rows] == [4, 5]


def test_missing_input_tab_is_an_error(conn, tmp_path):
    wb = Workbook()
    wb.active.title = "Something Else"
    path = tmp_path / "wrong.xlsx"
    wb.save(path)

    with pytest.raises(ValueError, match="Parts to lookup and where"):
        importer.import_parts(conn, path)
