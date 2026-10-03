"""Reader tests, driven by the type confusion in the client's real workbook.

2,930 part numbers in data/PN_List_full.xlsx are stored as numbers, 240 of
those in scientific notation, and 26 carry a date format on top. Every case
below corresponds to a real row.
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook

from txtav import xlsx_reader


class TestRenderNumber:
    @pytest.mark.parametrize(
        "raw,expected",
        [
            ("10003.0", "10003"),          # row 8236
            ("100040.0", "100040"),        # row 8238
            ("1.0005189E7", "10005189"),   # row 8240, scientific notation
            ("1.00057022E8", "100057022"), # row 8244
            ("3.5466514E7", "35466514"),   # row 74371, also date formatted
            ("2.4954017E7", "24954017"),   # row 8924
            ("0.0", "0"),
        ],
    )
    def test_whole_numbers_lose_the_decimal_tail(self, raw, expected):
        assert xlsx_reader.render_number(raw) == expected

    def test_fractional_values_keep_their_digits(self):
        assert xlsx_reader.render_number("1.25") == "1.25"

    def test_never_returns_scientific_notation(self):
        assert "E" not in xlsx_reader.render_number("1.00057022E8").upper()


class TestColumnIndex:
    @pytest.mark.parametrize(
        "ref,expected", [("A1", 0), ("B1", 1), ("Z9", 25), ("AA1", 26), ("AB100", 27)]
    )
    def test_maps_letters_to_zero_based_index(self, ref, expected):
        assert xlsx_reader.column_index(ref) == expected


@pytest.fixture
def mixed_workbook(tmp_path):
    """A workbook whose part number column mixes text and numeric cells."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Parts to lookup and where"
    ws.append(["PART NUMBER", "WH"])
    ws.append(["6622168-22", "TXT"])  # text
    ws.append([10003, "TXT"])          # stored as a number
    ws.append([100057022, "TXT"])      # large, would go scientific
    ws.append(["9912125-2EX", "TXT"])
    path = tmp_path / "mixed.xlsx"
    wb.save(path)
    return path


class TestIterRows:
    def test_numeric_part_numbers_arrive_without_a_decimal_tail(self, mixed_workbook):
        rows = dict(xlsx_reader.iter_rows(mixed_workbook, "Parts to lookup and where"))
        assert rows[3][0] == "10003"
        assert rows[4][0] == "100057022"

    def test_text_part_numbers_are_untouched(self, mixed_workbook):
        rows = dict(xlsx_reader.iter_rows(mixed_workbook, "Parts to lookup and where"))
        assert rows[2][0] == "6622168-22"
        assert rows[5][0] == "9912125-2EX"

    def test_second_column_is_read(self, mixed_workbook):
        rows = dict(xlsx_reader.iter_rows(mixed_workbook, "Parts to lookup and where"))
        assert all(values[1] == "TXT" for row, values in rows.items() if row > 1)

    def test_missing_sheet_names_what_is_available(self, mixed_workbook):
        with pytest.raises(xlsx_reader.WorkbookError, match="Parts to lookup and where"):
            list(xlsx_reader.iter_rows(mixed_workbook, "Nope"))
